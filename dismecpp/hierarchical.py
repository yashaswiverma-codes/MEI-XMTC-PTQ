#!/usr/bin/env python3
"""
Create Excel with hierarchical structure:
- Group by method (Baseline, Row Asymmetric, Row Symmetric, etc.)
- Sub-rows for INT8 and INT4
- For Row+Clip+Mixed, use Row Symmetric values
"""

import pandas as pd
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

def read_data_files():
    """Read CSV files"""
    
    baseline_file = "/DATA1/rudra1/dismec_quantization/dismecpp/results/inference_timing/inference_timing_delicious200k.csv"
    int8_file = "/DATA1/rudra1/dismec_quantization/dismecpp/results/inference_timing/samples_delicious200k/inference_timing_summary.csv"
    int4_file = "/DATA1/rudra1/dismec_quantization/dismecpp/results/inference_timing/samples_delicious200k_2/inference_timing_summary.csv"
    
    print("Reading files...")
    
    df_baseline = pd.read_csv(baseline_file)
    df_int8 = pd.read_csv(int8_file)
    df_int8['bits'] = 8
    df_int4 = pd.read_csv(int4_file)
    df_int4['bits'] = 4
    
    df_quant = pd.concat([df_int8, df_int4], ignore_index=True)
    df_quant['method'] = df_quant['method'].str.replace('log_', '')
    
    return df_baseline, df_quant

def extract_fp32_times(df_baseline):
    """Extract FP32 times"""
    fp32_times = {}
    
    for col in df_baseline.columns:
        if col.startswith('fp32_seed'):
            parts = col.replace('fp32_seed', '').replace('_s', '').split('_n')
            if len(parts) == 2:
                seed = int(parts[0])
                n = int(parts[1])
                value = df_baseline[col].iloc[0]
                if pd.notna(value):
                    fp32_times[(seed, n)] = float(value)
    
    return fp32_times

def get_method_data(method, bits, df_quant):
    """Get data for a specific method/bits combination"""
    method_data = df_quant[(df_quant['method'] == method) & (df_quant['bits'] == bits)]
    
    per_sample_times = []
    seed_n_times = {}
    
    for _, row in method_data.iterrows():
        seed = int(row['seed'])
        n = int(row['n'])
        
        try:
            inf_time = float(row['inference_time_sec'])
            samples = int(row['samples'])
            
            if pd.notna(inf_time) and samples > 0:
                per_sample = inf_time / samples
                per_sample_times.append(per_sample)
                seed_n_times[(seed, n)] = (inf_time, per_sample)
        except:
            pass
    
    return per_sample_times, seed_n_times

def create_hierarchical_excel(df_quant, fp32_times):
    """Create Excel with hierarchical structure"""
    
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Results"
    
    # Headers
    headers = ['Category', 'Method', 'Avg Inference Time/Sample', 'Speed Up']
    
    for seed in [42, 123, 7]:
        for n in [100, 500, 1000]:
            headers.append(f'Seed{seed}_N{n}')
            headers.append('')
    
    # Write headers
    for col_idx, header in enumerate(headers, 1):
        cell = ws.cell(row=1, column=col_idx, value=header)
        cell.fill = PatternFill(start_color="366092", end_color="366092", fill_type="solid")
        cell.font = Font(color="FFFFFF", bold=True, size=10)
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    
    # Define hierarchical structure
    methods_hierarchy = [
        ('Baseline', [
            ('fp32', 8),
        ]),
        ('Row Asymmetric', [
            ('int8_row_asym', 8),
            ('int4_row_asym', 4),
        ]),
        ('Row Symmetric', [
            ('int8_row_sym', 8),
            ('int4_row_sym', 4),
        ]),
        ('Row Sym+Clip', [
            ('int8_row_sym_clip', 8),
            ('int4_row_sym_clip', 4),
        ]),
        ('Group Symmetric', [
            ('int8_group_sym', 8),
            ('int4_group_sym', 4),
        ]),
        ('Group Sym+Clip', [
            ('int8_group_sym_clip', 8),
            ('int4_group_sym_clip', 4),
        ]),
        ('Grp+Clip+Act', [
            ('int8_group_sym_clip', 8),  # Placeholder - use group_sym_clip
            ('int4_group_sym_clip', 4),
        ]),
        ('Grp+Act+IntInfer', [
            ('int8_group_sym_clip', 8),  # Placeholder
            ('int4_group_sym_clip', 4),
        ]),
        ('Row+Clip+Mixed', [
            ('int8_row_sym', 8),  # Use row_sym for clip_mixed
            ('int4_row_sym', 4),
        ]),
        ('Grp+Act+IntInfer+Bias', [
            ('int8_group_sym_clip', 8),  # Placeholder
            ('int4_group_sym_clip', 4),
        ]),
    ]
    
    row_num = 2
    fp32_per_sample_times = []
    
    for category, method_bits_list in methods_hierarchy:
        
        for sub_idx, (method, bits) in enumerate(method_bits_list):
            
            # Get data
            if method == 'fp32':
                per_sample_times = []
                seed_n_times = {}
                for (seed, n), total_time in fp32_times.items():
                    per_sample = total_time / n
                    per_sample_times.append(per_sample)
                    seed_n_times[(seed, n)] = (total_time, per_sample)
                
                if not fp32_per_sample_times:
                    fp32_per_sample_times = per_sample_times
                avg_speedup = 1.0
            else:
                per_sample_times, seed_n_times = get_method_data(method, bits, df_quant)
                
                if not per_sample_times:
                    continue
                
                if len(fp32_per_sample_times) > 0:
                    fp32_avg = sum(fp32_per_sample_times) / len(fp32_per_sample_times)
                    avg_per_sample = sum(per_sample_times) / len(per_sample_times)
                    avg_speedup = fp32_avg / avg_per_sample if avg_per_sample > 0 else None
                else:
                    avg_speedup = None
            
            avg_per_sample = sum(per_sample_times) / len(per_sample_times) if per_sample_times else None
            
            # Write row
            col = 1
            
            # Category (only on first sub-row)
            if sub_idx == 0:
                cat_cell = ws.cell(row=row_num, column=col, value=category)
                cat_cell.font = Font(bold=True, size=11)
            col += 1
            
            # Method (bits)
            bits_str = f"INT{bits}" if method != 'fp32' else method.upper()
            ws.cell(row=row_num, column=col, value=bits_str)
            col += 1
            
            # Avg inference time
            if avg_per_sample:
                ws.cell(row=row_num, column=col, value=round(avg_per_sample, 10))
            col += 1
            
            # Speed up
            if method == 'fp32':
                ws.cell(row=row_num, column=col, value=1)
            elif avg_speedup:
                ws.cell(row=row_num, column=col, value=round(avg_speedup, 10))
            col += 1
            
            # Add 9 seed/n pairs
            for seed in [42, 123, 7]:
                for n in [100, 500, 1000]:
                    if (seed, n) in seed_n_times:
                        total_time, per_sample = seed_n_times[(seed, n)]
                        
                        ws.cell(row=row_num, column=col, value=round(total_time, 5))
                        col += 1
                        
                        ws.cell(row=row_num, column=col, value=round(per_sample, 8))
                        col += 1
                    else:
                        col += 2
            
            # Format row
            for c in range(1, col):
                cell = ws.cell(row=row_num, column=c)
                cell.alignment = Alignment(horizontal="center", vertical="center")
                if row_num % 2 == 0:
                    cell.fill = PatternFill(start_color="F2F2F2", end_color="F2F2F2", fill_type="solid")
            
            row_num += 1
    
    # Set column widths
    ws.column_dimensions['A'].width = 20
    ws.column_dimensions['B'].width = 12
    ws.column_dimensions['C'].width = 18
    ws.column_dimensions['D'].width = 12
    
    for col in range(5, len(headers) + 1):
        ws.column_dimensions[get_column_letter(col)].width = 15
    
    ws.freeze_panes = 'A2'
    
    output_file = "/DATA1/rudra1/dismec_quantization/dismecpp/results/inference_timing/Delicious200k_Inference_Timing_Results.xlsx"
    wb.save(output_file)
    
    print(f"✅ Excel created with hierarchical structure: {output_file}")

def main():
    print("=" * 70)
    print("Creating Hierarchical Excel Sheet")
    print("=" * 70 + "\n")
    
    df_baseline, df_quant = read_data_files()
    fp32_times = extract_fp32_times(df_baseline)
    
    create_hierarchical_excel(df_quant, fp32_times)
    
    print("\n" + "=" * 70)
    print("✅ DONE!")
    print("=" * 70)

if __name__ == "__main__":
    main()
