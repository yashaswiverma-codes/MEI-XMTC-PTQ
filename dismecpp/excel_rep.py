#!/usr/bin/env python3
"""
Create Excel sheet for Delicious200k inference timing results
Following the pattern: Method, Bits, Inference time per sample, Speed Up, Seed42_N100, etc.
For row_sym_clip_mixed where error occurred, use row_sym data instead
"""

import pandas as pd
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
import os

def read_and_organize_data():
    """Read CSV files and organize data by method and seed/n combinations"""
    
    # Read the two summary CSV files
    int8_file = "/DATA1/rudra1/dismec_quantization/dismecpp/results/inference_timing/samples_delicious200k/inference_timing_summary.csv"
    int4_file = "/DATA1/rudra1/dismec_quantization/dismecpp/results/inference_timing/samples_delicious200k_2/inference_timing_summary.csv"
    
    print("Reading data files...")
    
    # Read INT8 data
    df_int8 = pd.read_csv(int8_file)
    df_int8['bits'] = 8
    print(f"INT8 data: {len(df_int8)} rows")
    
    # Read INT4 data
    df_int4 = pd.read_csv(int4_file)
    df_int4['bits'] = 4
    print(f"INT4 data: {len(df_int4)} rows")
    
    # Combine
    df = pd.concat([df_int8, df_int4], ignore_index=True)
    
    # Remove 'log_' prefix from method names
    df['method'] = df['method'].str.replace('log_', '')
    
    print(f"Total records: {len(df)}")
    print(f"Methods: {sorted(df['method'].unique())}")
    
    return df

def calculate_per_sample_time(row):
    """Calculate inference time per sample"""
    if pd.isna(row['inference_time_sec']) or pd.isna(row['samples']):
        return None
    if row['samples'] == 0:
        return None
    return row['inference_time_sec'] / row['samples']

def get_baseline_time(df, seed, n):
    """Get FP32 baseline time for given seed and n"""
    baseline = df[(df['method'] == 'fp32') & (df['seed'] == seed) & (df['n'] == n)]
    if baseline.empty:
        return None
    per_sample = calculate_per_sample_time(baseline.iloc[0])
    return per_sample

def create_excel_sheet():
    """Create the Excel sheet with formatted data"""
    
    df = read_and_organize_data()
    
    # Calculate per-sample times
    df['per_sample_time'] = df.apply(calculate_per_sample_time, axis=1)
    
    # Get FP32 baseline per-sample times for all seed/n combinations
    baseline_times = {}
    for seed in [7, 42, 123]:
        for n in [100, 500, 1000]:
            baseline_times[(seed, n)] = get_baseline_time(df, seed, n)
    
    print("\nBaseline (FP32) per-sample times:")
    for (seed, n), time_val in baseline_times.items():
        print(f"  Seed{seed}_N{n}: {time_val}")
    
    # Organize data for Excel
    methods_order = [
        'fp32',
        'int8_row_asym',
        'int8_row_sym',
        'int8_row_sym_clip',
        'int8_row_sym_clip_mixed',
        'int8_group_sym',
        'int8_group_sym_clip',
        'int4_row_asym',
        'int4_row_sym',
        'int4_row_sym_clip',
        'int4_row_sym_clip_mixed',
        'int4_group_sym',
        'int4_group_sym_clip'
    ]
    
    # Create output data
    output_rows = []
    
    for method in methods_order:
        method_data = df[df['method'] == method]
        
        if method_data.empty:
            continue
        
        bits = method_data['bits'].iloc[0]
        
        # Get per-sample time (average across all seeds/n for this method)
        per_sample_times = method_data['per_sample_time'].dropna()
        if len(per_sample_times) > 0:
            avg_per_sample = per_sample_times.mean()
        else:
            avg_per_sample = None
        
        # Calculate speed up relative to FP32
        if method != 'fp32' and avg_per_sample is not None:
            fp32_baseline = df[(df['method'] == 'fp32')]['per_sample_time'].mean()
            if fp32_baseline is not None and fp32_baseline > 0:
                speed_up = fp32_baseline / avg_per_sample
            else:
                speed_up = None
        else:
            speed_up = 1.0 if method == 'fp32' else None
        
        row = {
            'Dataset': 'Delicious200k',
            'Method': method.replace('_', ' ').title(),
            'Bits': int(bits),
            'Inference Time/Sample': avg_per_sample,
            'Speed Up': speed_up
        }
        
        # Add seed/n combinations
        for seed in [42, 123, 7]:
            for n in [100, 500, 1000]:
                col_name = f'Seed{seed}_N{n}'
                seed_data = method_data[(method_data['seed'] == seed) & (method_data['n'] == n)]
                
                if not seed_data.empty:
                    per_sample = calculate_per_sample_time(seed_data.iloc[0])
                    baseline = baseline_times[(seed, n)]
                    
                    if method != 'fp32' and per_sample is not None and baseline is not None and baseline > 0:
                        speed_up_val = baseline / per_sample
                    elif method == 'fp32':
                        speed_up_val = 1.0
                    else:
                        speed_up_val = None
                    
                    # Store as tuple (time, speedup)
                    row[col_name] = (per_sample, speed_up_val)
                else:
                    # Data missing - for row_sym_clip_mixed, use row_sym instead
                    if 'clip_mixed' in method and seed == 42 and n == 100:
                        fallback_method = method.replace('_clip_mixed', '')
                        fallback_data = df[(df['method'] == fallback_method) & (df['seed'] == seed) & (df['n'] == n)]
                        if not fallback_data.empty:
                            per_sample = calculate_per_sample_time(fallback_data.iloc[0])
                            baseline = baseline_times[(seed, n)]
                            if per_sample is not None and baseline is not None and baseline > 0:
                                speed_up_val = baseline / per_sample
                            else:
                                speed_up_val = None
                            row[col_name] = (per_sample, speed_up_val)
                        else:
                            row[col_name] = None
                    else:
                        row[col_name] = None
        
        output_rows.append(row)
    
    # Create DataFrame
    result_df = pd.DataFrame(output_rows)
    
    # Save to Excel with formatting
    output_file = "/DATA1/rudra1/dismec_quantization/dismecpp/results/inference_timing/Delicious200k_Inference_Timing_Results.xlsx"
    
    with pd.ExcelWriter(output_file, engine='openpyxl') as writer:
        result_df.to_excel(writer, sheet_name='Results', index=False)
        
        workbook = writer.book
        worksheet = writer.sheets['Results']
        
        # Define styles
        header_fill = PatternFill(start_color="366092", end_color="366092", fill_type="solid")
        header_font = Font(color="FFFFFF", bold=True, size=11)
        
        method_fill = PatternFill(start_color="D9E1F2", end_color="D9E1F2", fill_type="solid")
        method_font = Font(bold=True)
        
        center_alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        thin_border = Border(
            left=Side(style='thin'),
            right=Side(style='thin'),
            top=Side(style='thin'),
            bottom=Side(style='thin')
        )
        
        # Format headers
        for col in range(1, len(result_df.columns) + 1):
            cell = worksheet.cell(row=1, column=col)
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = center_alignment
            cell.border = thin_border
        
        # Format data rows
        for row_idx in range(2, len(result_df) + 2):
            for col_idx in range(1, len(result_df.columns) + 1):
                cell = worksheet.cell(row=row_idx, column=col_idx)
                cell.border = thin_border
                cell.alignment = center_alignment
                
                # Alternate row colors for methods
                if row_idx % 2 == 0:
                    cell.fill = PatternFill(start_color="F2F2F2", end_color="F2F2F2", fill_type="solid")
        
        # Set column widths
        worksheet.column_dimensions['A'].width = 15
        worksheet.column_dimensions['B'].width = 25
        worksheet.column_dimensions['C'].width = 8
        
        for col_idx in range(4, len(result_df.columns) + 1):
            worksheet.column_dimensions[get_column_letter(col_idx)].width = 20
        
        # Freeze header row
        worksheet.freeze_panes = 'A2'
    
    print(f"\n✅ Excel file created: {output_file}")
    print(f"   Total rows: {len(result_df)}")
    print(f"   Columns: {list(result_df.columns)}")
    
    return output_file

if __name__ == "__main__":
    create_excel_sheet()
